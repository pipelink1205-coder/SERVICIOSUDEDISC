"""Exportación a Excel de agregados (nunca microdato)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .ingesta.lector import MESES

ETIQUETAS = {
    "servicio": "Servicios", "indicador_cod": "Indicadores", "proyecto_cod": "Proyectos", "sexo": "Sexo",
    "discapacidad": "Discapacidad", "estrato": "Estrato", "escolaridad": "Escolaridad",
    "salud": "Régimen de salud", "etnia": "Etnia", "pais": "País de nacimiento", "area": "Área",
    "rango_edad": "Rango de edad",
}


def libro_agregados(t: dict, corte: dict, filtros_txt: str) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        k = t["kpis"]
        resumen = pd.DataFrame([
            ("Corte", f"{MESES[corte['mes_corte'] - 1].capitalize()} {corte['anio']}"),
            ("Archivo", corte["archivo"]),
            ("Filtros", filtros_txt or "Sin filtros"),
            ("Atenciones", k["atenciones"]), ("Personas distintas", k["personas"]),
            ("Mujeres (personas)", k["mujeres"]["personas"]), ("Hombres (personas)", k["hombres"]["personas"]),
            ("Edad promedio por persona", k["edad_promedio"]),
            ("Urbana (atenciones)", k["urbana"]["atenciones"]), ("Rural (atenciones)", k["rural"]["atenciones"]),
        ], columns=["Indicador", "Valor"])
        resumen.to_excel(xw, sheet_name="Resumen", index=False)
        pd.DataFrame(t["serie"]).rename(columns={
            "mes": "Mes", "nombre": "Nombre", "atenciones": "Atenciones", "personas": "Personas",
            "personas_primera_vez": "Personas por primera vez"}).to_excel(xw, sheet_name="Serie mensual", index=False)
        terr = t["territorio"]["territorios"] + t["territorio"]["sin_mapa"]
        pd.DataFrame(terr).rename(columns={"numero": "N.º", "nombre": "Comuna / corregimiento",
                                           "atenciones": "Atenciones", "personas": "Personas"}) \
            .to_excel(xw, sheet_name="Territorio", index=False)
        for clave, hoja in ETIQUETAS.items():
            filas = t["dimensiones"].get(clave, [])
            if filas:
                pd.DataFrame(filas).rename(columns={"valor": hoja, "atenciones": "Atenciones",
                                                    "personas": "Personas", "indicador": "Indicador",
                                                    "nombre": "Nombre"}) \
                    .to_excel(xw, sheet_name=hoja[:31], index=False)
        pd.DataFrame(t["piramide"]).rename(columns={"rango": "Rango", "hombres": "Hombres", "mujeres": "Mujeres",
                                                    "otro": "Otro"}).to_excel(xw, sheet_name="Pirámide", index=False)

        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1E3A8A")
                c.alignment = Alignment(vertical="center")
            for i, col in enumerate(ws.columns, 1):
                ancho = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 70)
            ws.freeze_panes = "A2"
    return buf.getvalue()
