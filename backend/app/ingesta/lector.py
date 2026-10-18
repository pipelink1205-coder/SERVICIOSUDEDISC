"""Lectura del libro mensual y conversión a atenciones sin datos personales.

Reglas:
- Se aceptan encabezados con o sin asterisco, saltos de línea, tildes o espacios dobles.
- Una fila es atención válida solo si el mes de reporte es un mes legible.
- El documento se convierte en un código irreversible (sha256 con sal local). Nombres,
  dirección, fecha de nacimiento, orientación, identidad de género y observaciones
  NO se guardan.
- Los valores no se renombran en silencio: las variantes se reportan en calidad.
"""
from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ..config import HOJA_PRODUCTO, SAL_DOCUMENTO

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
MES_NUM = {m: i + 1 for i, m in enumerate(MESES)}
MES_NUM["SETIEMBRE"] = 9

RANGOS_EDAD = ["0–12", "13–17", "18–29", "30–44", "45–59", "60+"]
TERRITORIOS_MAPA = set(range(1, 17)) | {50, 60, 70, 80, 90}
VALORES_VACIOS = {"", "NAN", "NONE", "NULL", "N/A", "NA"}


class ErrorLectura(Exception):
    """Error que se muestra tal cual al usuario."""


def quitar_tildes(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar_encabezado(texto: object) -> str:
    t = quitar_tildes(str(texto or "")).upper()
    t = t.replace("*", " ").replace('"', " ")
    return re.sub(r"\s+", " ", t).strip()


# (prefijo normalizado, campo interno). El orden importa: el primero que coincide gana.
COLUMNAS = [
    ("MES DE REPORTE", "mes"),
    ("CODIGO PROYECTO", "proyecto_cod"),
    ("NOMBRE DEL PROYECTO", "proyecto"),
    ("CODIGO INDICADOR", "indicador_cod"),
    ("NOMBRE INDICADOR", "indicador"),
    ("CODIGO VALOR ESTADISTICO", "codigo_valor"),
    ("TIPO DOCUMENTO IDENTIDAD", "tipo_doc"),
    ("NUMERO DOCUMENTO IDENTIDAD", "documento"),
    ("SEXO", "sexo"),
    ("ETNIA", "etnia"),
    ("CONDICION DE DISCAPACIDAD", "discapacidad"),
    ("CONDICION DE VICTIMA", "victima"),
    ("HABITANZA EN CALLE", "habitanza"),
    ("CAMPESINO", "campesino"),
    ("ESCOLARIDAD", "escolaridad"),
    ("REGIMEN DE AFILIACION", "salud"),
    ("ESTRATO", "estrato"),
    ("AREA", "area"),
    ("BARRIO", "barrio"),
    ("COMUNA", "comuna_txt"),
    ("FECHA INGRESO", "fecha_ingreso"),
    ("ANOS CUMPLIDOS", "edad"),
    ("ANOS DE PERMANENCIA", "perm_anios"),
    ("MESES DE PERMANENCIA", "perm_meses"),
    ("NUEVO", "nuevo"),
    ("VIENE", "viene"),
    ("PAIS DE NACIMIENTO", "pais"),
    ("BIEN, PRODUCTO, SERVICIO", "servicio"),
]
OBLIGATORIAS = {"mes": "MES DE REPORTE", "servicio": "BIEN, PRODUCTO, SERVICIO RECIBIDO",
                "documento": "NÚMERO DOCUMENTO IDENTIDAD", "sexo": "* SEXO"}

# Campos categóricos que se revisan en calidad (vacíos, SIN DATO, variantes).
DIMENSIONES = {
    "servicio": "Servicio", "indicador_cod": "Indicador", "sexo": "Sexo", "etnia": "Etnia",
    "discapacidad": "Discapacidad", "victima": "Hecho victimizante", "habitanza": "Habitanza en calle",
    "campesino": "Campesino", "escolaridad": "Escolaridad", "salud": "Régimen de salud",
    "estrato": "Estrato", "area": "Área", "comuna": "Comuna", "barrio": "Barrio",
    "pais": "País de nacimiento", "tipo_doc": "Tipo de documento",
}


@dataclass
class Resultado:
    hoja: str
    hojas_libro: list[str]
    filas_leidas: int
    atenciones: pd.DataFrame
    calidad: list[dict] = field(default_factory=list)
    anio: int | None = None
    mes_corte: int | None = None


# ---------------------------------------------------------------------------
def hash_archivo(ruta: Path) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def inferir_periodo(nombre: str) -> tuple[int | None, int | None]:
    """'Reporte PI diciembre_2025 Discapacidad.xlsm' -> (2025, 12)."""
    n = quitar_tildes(nombre).upper()
    anio = re.search(r"(20\d{2})", n)
    mes = next((MES_NUM[m] for m in MES_NUM if re.search(rf"(^|[^A-Z]){m}([^A-Z]|$)", n)), None)
    return (int(anio.group(1)) if anio else None), mes


def detectar_hoja(ruta: Path) -> tuple[str, list[str]]:
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hojas = list(wb.sheetnames)
    finally:
        wb.close()
    objetivo = {normalizar_encabezado(h).replace(".", "").replace(" ", "") for h in HOJA_PRODUCTO}
    for h in hojas:
        if normalizar_encabezado(h).replace(".", "").replace(" ", "") in objetivo:
            return h, hojas
    raise ErrorLectura(
        "No se encontró la hoja 'ESTAD. PERSONAS IND. PRODUCTO'. Hojas del libro: " + ", ".join(hojas)
    )


def _limpiar(serie: pd.Series) -> pd.Series:
    s = serie.astype("string").str.replace(r"\s+", " ", regex=True).str.strip().str.upper()
    return s.where(~s.isin(VALORES_VACIOS) & s.notna(), None)


def _mes_a_numero(valor: object) -> int | None:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return None
    t = quitar_tildes(str(valor)).upper().strip()
    if not t:
        return None
    for nombre, num in MES_NUM.items():
        if re.fullmatch(rf"(\d{{1,2}}\s*-?\s*)?{nombre}", t):
            return num
    if re.fullmatch(r"\d{1,2}(\.0)?", t):
        n = int(float(t))
        return n if 1 <= n <= 12 else None
    fecha = pd.to_datetime(t, errors="coerce")
    if pd.notna(fecha) and re.match(r"\d{4}-\d{2}-\d{2}", t):
        return int(fecha.month)
    return None


def _codigo_persona(doc: object) -> str | None:
    if doc is None or pd.isna(doc):
        return None
    digitos = re.sub(r"\.0$", "", str(doc).strip())
    digitos = re.sub(r"[^0-9A-Za-z]", "", digitos).upper()
    if not digitos:
        return None
    return hashlib.sha256((SAL_DOCUMENTO + digitos).encode()).hexdigest()[:20]


def _marca_x(serie: pd.Series) -> pd.Series:
    return serie.astype("string").str.strip().str.upper().fillna("").eq("X").astype(int)


def _filas_texto(filas: pd.Series | list, maximo: int = 20) -> str:
    lista = list(filas)[:maximo]
    return ", ".join(str(int(f)) for f in lista)


# ---------------------------------------------------------------------------
def leer_libro(ruta: Path, nombre_archivo: str | None = None) -> Resultado:
    ruta = Path(ruta)
    if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ErrorLectura(f"El archivo debe ser .xlsx o .xlsm (llegó {ruta.suffix or 'sin extensión'}).")

    hoja, hojas = detectar_hoja(ruta)
    crudo = pd.read_excel(ruta, sheet_name=hoja, dtype=object, engine="openpyxl")
    filas_leidas = len(crudo)

    # --- mapear columnas -----------------------------------------------------
    renombre: dict[str, str] = {}
    usados: set[str] = set()
    for col in crudo.columns:
        n = normalizar_encabezado(col)
        for prefijo, campo in COLUMNAS:
            if campo not in usados and n.startswith(prefijo):
                renombre[col] = campo
                usados.add(campo)
                break
    faltan = [v for k, v in OBLIGATORIAS.items() if k not in usados]
    if faltan:
        raise ErrorLectura("Faltan columnas obligatorias en la hoja: " + ", ".join(faltan))

    df = crudo.rename(columns=renombre)[list(renombre.values())].copy()
    df["fila"] = df.index + 2                          # fila real en Excel (encabezado en 1)

    calidad: list[dict] = []

    # --- filas no válidas (mes ilegible) ---------------------------------------
    df["mes_num"] = df["mes"].map(_mes_a_numero)
    invalidas = df[df["mes_num"].isna()]
    for valor, grupo in invalidas.groupby(invalidas["mes"].astype("string").fillna("(vacío)"), dropna=False):
        calidad.append(dict(
            tipo="rechazo", campo="Mes de reporte", valor=str(valor), filas=len(grupo),
            detalle="El mes no es legible: la fila no es una atención (encabezado repetido o código).",
            ejemplo_filas=_filas_texto(grupo["fila"]),
        ))
    df = df[df["mes_num"].notna()].copy()
    df["mes"] = df["mes_num"].astype(int)

    # --- limpiar categóricos ----------------------------------------------------
    for c in ["proyecto_cod", "proyecto", "indicador_cod", "indicador", "servicio", "tipo_doc", "sexo",
              "etnia", "discapacidad", "victima", "habitanza", "campesino", "escolaridad", "salud",
              "estrato", "area", "barrio", "comuna_txt", "pais"]:
        if c not in df:
            df[c] = None
        df[c] = _limpiar(df[c])
    df["proyecto_cod"] = df["proyecto_cod"].str.replace(r"\.0$", "", regex=True)
    # el nombre del indicador conserva mayúsculas/minúsculas originales
    if "indicador" in renombre.values():
        df["indicador"] = crudo.loc[df.index, [k for k, v in renombre.items() if v == "indicador"][0]] \
            .astype("string").str.strip()

    # --- territorio -----------------------------------------------------------------
    num = df["comuna_txt"].str.extract(r"^\s*(\d{1,2})")[0]
    df["comuna_num"] = pd.to_numeric(num, errors="coerce").astype("Int64")
    df["comuna"] = df["comuna_txt"]

    # --- edad y permanencia -------------------------------------------------------
    df["edad"] = pd.to_numeric(df.get("edad"), errors="coerce").round().astype("Int64")
    edad_ok = df["edad"].between(0, 110)
    df.loc[~edad_ok.fillna(False), "edad"] = pd.NA
    df["rango_edad"] = pd.cut(df["edad"].astype("float"), bins=[-1, 12, 17, 29, 44, 59, 200],
                              labels=RANGOS_EDAD).astype("string")
    pa = pd.to_numeric(df.get("perm_anios"), errors="coerce").fillna(0)
    pm = pd.to_numeric(df.get("perm_meses"), errors="coerce").fillna(0)
    df["permanencia_meses"] = (pa * 12 + pm).astype(int)
    fi = pd.to_datetime(df.get("fecha_ingreso"), errors="coerce")
    df["fecha_ingreso"] = fi.dt.strftime("%Y-%m-%d")
    df["nuevo"] = _marca_x(df["nuevo"]) if "nuevo" in df else 0
    df["viene"] = _marca_x(df["viene"]) if "viene" in df else 0

    # --- identidad --------------------------------------------------------------------
    df["persona"] = df["documento"].map(_codigo_persona)

    # --- periodo del corte ----------------------------------------------------------------
    anio, mes_corte = inferir_periodo(nombre_archivo or ruta.name)
    if mes_corte is None:
        mes_corte = int(df["mes"].max()) if len(df) else None
    if anio is None and fi.notna().any():
        anio = int(fi.dt.year.max())

    # --- calidad -------------------------------------------------------------------------
    calidad += _revisar_calidad(df, edad_ok)

    columnas = ["fila", "mes", "proyecto_cod", "proyecto", "indicador_cod", "indicador", "servicio",
                "persona", "tipo_doc", "sexo", "edad", "rango_edad", "etnia", "discapacidad", "victima",
                "habitanza", "campesino", "escolaridad", "salud", "estrato", "area", "comuna_num",
                "comuna", "barrio", "fecha_ingreso", "permanencia_meses", "nuevo", "viene", "pais"]
    atenciones = df[columnas].reset_index(drop=True)
    return Resultado(hoja=hoja, hojas_libro=hojas, filas_leidas=filas_leidas, atenciones=atenciones,
                     calidad=calidad, anio=anio, mes_corte=mes_corte)


def _clave_variante(v: str) -> str:
    v = quitar_tildes(v).upper()
    v = re.sub(r"\b(DE|DEL|EN|LA|EL|LOS|LAS)\b", " ", v)
    v = re.sub(r"[^A-Z0-9]", "", v)
    return re.sub(r"(O|A)$", "", v) if len(v) > 4 else v      # NINGUNA ~ NINGUNO


def _revisar_calidad(df: pd.DataFrame, edad_ok: pd.Series) -> list[dict]:
    q: list[dict] = []

    for campo, etiqueta in DIMENSIONES.items():
        s = df[campo]
        vacio = s.isna()
        if vacio.any():
            q.append(dict(tipo="vacio", campo=etiqueta, valor="(vacío)", filas=int(vacio.sum()),
                          detalle="Celda vacía.", ejemplo_filas=_filas_texto(df.loc[vacio, "fila"])))
        sin = s.eq("SIN DATO").fillna(False)
        if sin.any():
            q.append(dict(tipo="sin_dato", campo=etiqueta, valor="SIN DATO", filas=int(sin.sum()),
                          detalle="Registrado explícitamente como SIN DATO.",
                          ejemplo_filas=_filas_texto(df.loc[sin, "fila"])))
        # variantes de escritura del mismo valor
        valores = s.dropna().value_counts()
        grupos: dict[str, list[str]] = {}
        for v in valores.index:
            grupos.setdefault(_clave_variante(v), []).append(v)
        for vs in grupos.values():
            if len(vs) > 1:
                total = int(sum(valores[v] for v in vs))
                q.append(dict(tipo="variante", campo=etiqueta, valor=" | ".join(vs), filas=total,
                              detalle="Mismo valor escrito de formas distintas: "
                                      + "; ".join(f"{v} ({int(valores[v])})" for v in vs),
                              ejemplo_filas=""))

    # documento vacío
    sin_doc = df["persona"].isna()
    if sin_doc.any():
        q.append(dict(tipo="vacio", campo="Documento", valor="(vacío)", filas=int(sin_doc.sum()),
                      detalle="Cuenta como atención pero no como persona identificada.",
                      ejemplo_filas=_filas_texto(df.loc[sin_doc, "fila"])))

    # edad fuera de rango o ilegible
    mala = ~edad_ok.fillna(False)
    if mala.any():
        q.append(dict(tipo="inconsistencia", campo="Edad", valor="fuera de 0–110 o vacía",
                      filas=int(mala.sum()), detalle="No entra en la pirámide ni en la edad promedio.",
                      ejemplo_filas=_filas_texto(df.loc[mala, "fila"])))

    # territorio sin polígono
    fuera = df["comuna_num"].isna() | ~df["comuna_num"].isin(list(TERRITORIOS_MAPA))
    if fuera.any():
        for valor, g in df[fuera].groupby(df.loc[fuera, "comuna"].fillna("(vacío)")):
            q.append(dict(tipo="fuera_mapa", campo="Comuna", valor=str(valor), filas=len(g),
                          detalle="No tiene polígono en el mapa: se muestra en la tabla de no georreferenciados.",
                          ejemplo_filas=_filas_texto(g["fila"])))

    # duplicados: misma persona, mismo servicio, mismo mes
    con_doc = df[df["persona"].notna()]
    dup = con_doc.duplicated(subset=["persona", "servicio", "mes"], keep="first")
    if dup.any():
        q.append(dict(tipo="duplicado", campo="Persona + servicio + mes", valor="repetido",
                      filas=int(dup.sum()),
                      detalle="La misma persona aparece más de una vez en el mismo servicio y mes.",
                      ejemplo_filas=_filas_texto(con_doc.loc[dup, "fila"])))

    # una persona con más de un sexo registrado
    sexos = con_doc.dropna(subset=["sexo"]).groupby("persona")["sexo"].nunique()
    varias = sexos[sexos > 1]
    if len(varias):
        filas = con_doc[con_doc["persona"].isin(varias.index)]["fila"]
        q.append(dict(tipo="inconsistencia", campo="Sexo", valor="más de un valor por persona",
                      filas=int(len(varias)), detalle="Personas con sexo distinto entre registros.",
                      ejemplo_filas=_filas_texto(filas)))
    return q
